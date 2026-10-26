import json
from io import BytesIO
from typing import Annotated, Optional

from fastapi import APIRouter, Depends, File, Form, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.dependencies import get_current_user, require_role
from app.models.contact import Contact
from app.models.deal import Deal
from app.models.user import User
from app.schemas.activity import ActivityCreate
from app.schemas.contact import (
    ArchiveRequest,
    ContactCreate,
    ContactUpdate,
    DuplicateCheckRequest,
    DuplicateResultResponse,
)
from app.services import (
    access_service,
    activity_service,
    autocount_service,
    chatbot_send_service,
    contact_service,
    duplicate_service,
    routing_service,
)
from app.utils.response import fail, ok

router = APIRouter()


def _not_found():
    # Also the answer for a customer that is not yours, so an id's existence is
    # never confirmed to someone who may not see it.
    return fail("Contact not found", code="NOT_FOUND", status_code=404)


# The fields a `column_map` may name (docs/contact-import-columns.md section 2).
IMPORT_FIELDS = (
    "name",
    "company",
    "industry",
    "email",
    "phone",
    "address",
    "notes",
    "tags",
    "assigned_to_email",
    "status",
    "priority",
    "deal_value",
    "deal_title",
    "id",
)

# Exactly the nine characters the frontend strips too (design section 9, item 16);
# Python's own str.strip() set differs from JavaScript's trim().
_HEADER_WHITESPACE = " \t\n\r\x0b\x0c 　﻿"


class _JsonPairs(list):
    """A JSON object's key/value pairs, kept so that repeated keys can be seen."""


def _parse_column_map(raw: str) -> tuple[Optional[dict[str, Optional[str]]], Optional[str]]:
    """`column_map` keyed by stripped header, or the reason it is refused."""
    try:
        parsed = json.loads(raw, object_pairs_hook=_JsonPairs)
    except ValueError:
        return None, "column_map is not valid JSON"
    if not isinstance(parsed, _JsonPairs):
        return None, "column_map must be a JSON object"

    mapping: dict[str, Optional[str]] = {}
    for key, value in parsed:
        header = key.strip(_HEADER_WHITESPACE)
        if not header:
            return None, "column_map has an empty header"
        if header in mapping:
            return None, f"column_map names the header more than once: {header}"
        if value is not None and not (isinstance(value, str) and value in IMPORT_FIELDS):
            return None, f"column_map has an unknown field for {header}"
        mapping[header] = value

    fields = [f for f in mapping.values() if f is not None]
    if len(fields) != len(set(fields)):
        return None, "column_map maps the same field to more than one column"
    if "name" not in fields and "id" not in fields:
        return None, "column_map must map a column to name or id"
    return mapping, None


def _text_header_columns(ws) -> dict[str, int]:
    """Stripped text header -> index of the first column carrying it.

    Numbers, dates, formulas and blank cells are left out: the browser reads them
    differently from openpyxl, so a `column_map` could never match them reliably.
    """
    columns: dict[str, int] = {}
    for idx, cell in enumerate(next(ws.iter_rows(max_row=1), ())):
        if cell.data_type != "s" or not isinstance(cell.value, str):
            continue
        header = cell.value.strip(_HEADER_WHITESPACE)
        if header and header not in columns:
            columns[header] = idx
    return columns


@router.get("")
async def list_contacts(
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(get_current_user)],
    search: Optional[str] = Query(None),
    industry: Optional[str] = Query(None),
    assigned_to: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=500),
    sort_by: str = Query("created_at"),
    order: str = Query("desc"),
    is_archived: int = Query(0),
    unassigned: int = Query(0, ge=0, le=1),
):
    result = await contact_service.list_contacts(
        db,
        current_user,
        search,
        industry,
        assigned_to,
        page,
        page_size,
        sort_by,
        order,
        is_archived,
        unassigned,
    )
    return ok(data=result)


@router.get("/import/template")
async def download_import_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "customer template"
    headers = [
        "name",
        "company",
        "industry",
        "email",
        "phone",
        "address",
        "remark",
        "tag",
        "assigned_to_email",
        "status",
        "priority",
        "amount",
        "deal_title",
        "customer_id (add deal only)",
    ]
    ws.append(headers)
    # Example row — new contact
    ws.append(
        [
            "Zhang San",
            "Example Tech Co.",
            "Technology/IT",
            "zhangsan@example.com",
            "13800138000",
            "123 Main St, Beijing",
            "Example note",
            "VIP,key account",
            "sales@crm.com",
            "lead",
            "mid",
            "100000",
            "",
            "",
        ]
    )
    # Example row — add deal to existing contact
    ws.append(
        [
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "negotiating",
            "high",
            "50000",
            "Q2 renewal",
            "existing-contact-uuid",
        ]
    )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="contact_import_template.xlsx"'},
    )


@router.post("/import")
async def import_contacts(
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(get_current_user)],
    file: UploadFile = File(...),
    auto_assign: bool = Form(False),
    assign_strategy: str = Form("rules"),
    column_map: Optional[str] = Form(None),
    skip_duplicates: bool = Form(True),
):
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        return fail("Only .xlsx or .xls files are supported", status_code=400)

    # "me" is resolved by the service from current_user; the rest are routing strategies.
    if assign_strategy not in (*routing_service.ASSIGN_STRATEGIES, "me"):
        return fail(f"Unknown assign_strategy: {assign_strategy}", status_code=400)

    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        return fail("File size cannot exceed 10MB", status_code=400)

    wb = load_workbook(BytesIO(content), read_only=True)
    ws = wb.active

    field_map = {
        "name": "name",
        "company": "company",
        "industry": "industry",
        "email": "email",
        "phone": "phone",
        "address": "address",
        "remark": "notes",
        "tag": "tags",
        "assigned_to_email": "assigned_to_email",
        "status": "status",
        "priority": "priority",
        "amount": "deal_value",
        "deal_title": "deal_title",
        "customer_id (add deal only)": "id",
    }

    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        return fail("File is empty", status_code=400)

    col_map = {}
    if column_map:
        # Only the columns the map names are read; anything else in the sheet,
        # template header or not, is left out.
        mapping, error = _parse_column_map(column_map)
        if error is None:
            header_columns = _text_header_columns(ws)
            missing = [h for h in mapping if h not in header_columns]
            if missing:
                error = f"column_map names a header not in the sheet: {missing[0]}"
        if error is not None:
            wb.close()
            return fail(error, code="INVALID_COLUMN_MAP", status_code=422)
        for h, field_name in mapping.items():
            if field_name is not None:
                col_map[header_columns[h]] = field_name
    else:
        # Map headers
        for idx, cell in enumerate(header):
            cell_str = str(cell or "").strip()
            if cell_str in field_map:
                col_map[idx] = field_map[cell_str]

    parsed_rows = []
    for row in rows_iter:
        row_dict = {}
        for idx, field_name in col_map.items():
            val = row[idx] if idx < len(row) else None
            row_dict[field_name] = str(val).strip() if val is not None else ""
        if any(v for v in row_dict.values()):
            parsed_rows.append(row_dict)

    wb.close()

    result = await contact_service.import_contacts(
        db,
        parsed_rows,
        current_user,
        auto_assign=auto_assign,
        assign_strategy=assign_strategy,
        skip_duplicates=skip_duplicates,
    )
    return ok(data=result)


# --- Duplicate routes (docs/contact-duplicate-hints.md 6.2 and 9.1) ---
# Declared before the /{contact_id} routes so "duplicate-check" is never read as an id.


def _duplicates_payload(result: duplicate_service.DuplicateResult) -> dict:
    return DuplicateResultResponse.model_validate(result).model_dump()


@router.post("/duplicate-check")
async def check_duplicates(
    body: DuplicateCheckRequest,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(get_current_user)],
):
    if not any((value or "").strip() for value in body.model_dump().values()):
        return ok(data=DuplicateResultResponse().model_dump())
    result = await duplicate_service.find_duplicates(
        db,
        user=current_user,
        name=body.name,
        company=body.company,
        email=body.email,
        phone=body.phone,
    )
    return ok(data=_duplicates_payload(result))


@router.get("/{contact_id}/duplicates")
async def list_duplicates(
    contact_id: str,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(get_current_user)],
):
    if not await access_service.may_access_contact(db, current_user, contact_id):
        return _not_found()
    # may_access_contact lets an admin through for any id, so missing and
    # deleted customers are refused here.
    contact = (
        await db.execute(
            select(Contact).where(Contact.id == contact_id, Contact.deleted_at.is_(None))
        )
    ).scalar_one_or_none()
    if contact is None:
        return _not_found()
    result = await duplicate_service.find_duplicates(
        db,
        user=current_user,
        name=contact.name,
        company=contact.company,
        email=contact.email,
        phone=contact.phone,
        exclude_id=contact.id,
    )
    return ok(data=_duplicates_payload(result))


@router.get("/{contact_id}")
async def get_contact(
    contact_id: str,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(get_current_user)],
):
    if not await access_service.may_access_contact(db, current_user, contact_id):
        return _not_found()
    # The user decides whether phone_conflict is worked out (admins only).
    contact = await contact_service.get_contact(db, contact_id, current_user)
    if not contact:
        return _not_found()
    return ok(data=contact)


@router.post("/{contact_id}/whatsapp/handback")
async def handback_whatsapp(
    contact_id: str,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(get_current_user)],
):
    if not await access_service.may_access_contact(db, current_user, contact_id):
        return fail("This customer is not assigned to you.", code="NOT_ASSIGNED", status_code=403)
    result = await chatbot_send_service.handback_to_ai(
        db, contact_id=contact_id, agent_id=current_user.id, agent_name=current_user.name
    )
    if result.category == chatbot_send_service.NOT_FOUND:
        return fail("This customer is not assigned to you.", code="NOT_ASSIGNED", status_code=403)
    if result.category == chatbot_send_service.NO_IDENTITY:
        return fail("Contact has no WhatsApp identity", code="NO_IDENTITY", status_code=400)
    if result.category == chatbot_send_service.NOT_CONFIGURED:
        return fail("Chatbot is not configured", code="NOT_CONFIGURED", status_code=503)
    if result.category == chatbot_send_service.CHATBOT_FAILED:
        return fail("Chatbot handback failed", code="CHATBOT_FAILED", status_code=502)
    return ok(data=contact_service.handover_to_dict(result.contact))


@router.post("")
async def create_contact(
    body: ContactCreate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(get_current_user)],
):
    data = body.model_dump()
    # Same rule as handing over an existing customer in update_contact.
    if data.get("assigned_to") and not await access_service.may_assign_to(
        db, current_user, data["assigned_to"]
    ):
        return fail("Permission denied", code="FORBIDDEN", status_code=403)
    contact = await contact_service.create_contact(db, data, current_user)
    return ok(data=contact, message="Contact created", status_code=201)


@router.put("/{contact_id}")
async def update_contact(
    contact_id: str,
    body: ContactUpdate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(get_current_user)],
):
    if not await access_service.may_access_contact(db, current_user, contact_id):
        return _not_found()
    data = {k: v for k, v in body.model_dump().items() if v is not None}
    if "is_archived" in data and current_user.role not in ("admin", "manager"):
        return fail("Permission denied", code="FORBIDDEN", status_code=403)
    # Resubmitting the current owner passes: having reached the customer, the
    # caller can always reach whoever owns it.
    if data.get("assigned_to") and not await access_service.may_assign_to(
        db, current_user, data["assigned_to"]
    ):
        return fail("Permission denied", code="FORBIDDEN", status_code=403)
    contact = await contact_service.update_contact(db, contact_id, data, current_user)
    if not contact:
        return _not_found()
    return ok(data=contact, message="Updated successfully")


@router.patch("/{contact_id}/archive")
async def archive_contact(
    contact_id: str,
    body: ArchiveRequest,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(require_role("admin", "manager"))],
):
    if not await access_service.may_access_contact(db, current_user, contact_id):
        return _not_found()
    contact = await contact_service.archive_contact(db, contact_id, body.is_archived)
    if not contact:
        return _not_found()
    msg = "Archived" if body.is_archived else "Unarchived"
    return ok(data=contact, message=f"{msg} successfully")


@router.delete("/{contact_id}")
async def delete_contact(
    contact_id: str,
    db: Annotated[AsyncSession, Depends(get_db)],
    _current_user: Annotated[User, Depends(require_role("admin"))],
):
    deleted = await contact_service.soft_delete_contact(db, contact_id)
    if not deleted:
        return fail("Contact not found", code="NOT_FOUND", status_code=404)
    return ok(message="Deleted successfully")


# --- AutoCount routes ---


@router.get("/{contact_id}/autocount-documents")
async def list_autocount_documents(
    contact_id: str,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(get_current_user)],
):
    if not await access_service.may_access_contact(db, current_user, contact_id):
        return _not_found()
    documents = await autocount_service.list_contact_documents(db, contact_id)
    return ok(data=documents)


# --- Activity routes ---


@router.get("/{contact_id}/activities")
async def list_activities(
    contact_id: str,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(get_current_user)],
):
    if not await access_service.may_access_contact(db, current_user, contact_id):
        return _not_found()
    activities = await activity_service.list_activities(db, contact_id)
    return ok(data=activities)


@router.post("/{contact_id}/activities")
async def create_activity(
    contact_id: str,
    body: ActivityCreate,
    db: Annotated[AsyncSession, Depends(get_db)],
    current_user: Annotated[User, Depends(get_current_user)],
):
    if not await access_service.may_access_contact(db, current_user, contact_id):
        return _not_found()
    # The deal has to be one of this customer's, or the note lands on a card
    # that belongs to somebody else.
    deal_matches = (
        await db.execute(
            select(Deal.id).where(
                Deal.id == body.deal_id,
                Deal.contact_id == contact_id,
                Deal.deleted_at.is_(None),
            )
        )
    ).scalar_one_or_none()
    if not deal_matches:
        return fail("Deal not found", code="NOT_FOUND", status_code=404)
    activity = await activity_service.create_activity(
        db,
        contact_id,
        body.deal_id,
        current_user.id,
        body.type,
        body.content,
        body.follow_date,
    )
    return ok(data=activity, message="Follow-up activity recorded", status_code=201)
